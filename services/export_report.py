"""Template-based report and tracker exports from confirmed findings only."""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from sqlalchemy.orm import Session, joinedload

from app.models import AnalystDecision, DecisionValue, Engagement, FindingGroup


def _confirmed_groups(db: Session, engagement_id: int) -> list[FindingGroup]:
    return (
        db.query(FindingGroup)
        .options(joinedload(FindingGroup.asset), joinedload(FindingGroup.decisions))
        .filter(
            FindingGroup.engagement_id == engagement_id,
            FindingGroup.queue_status == DecisionValue.CONFIRMED,
        )
        .order_by(FindingGroup.priority_score.desc(), FindingGroup.id)
        .all()
    )


def _latest_decision(group: FindingGroup) -> AnalystDecision | None:
    if not group.decisions:
        return None
    return sorted(group.decisions, key=lambda d: d.created_at, reverse=True)[0]


def export_tracker_csv(db: Session, engagement_id: int) -> bytes:
    groups = _confirmed_groups(db, engagement_id)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        [
            "finding_group_id",
            "title",
            "severity",
            "tool",
            "rule_id",
            "hostname",
            "ip_address",
            "port",
            "protocol",
            "decision",
            "decision_reason",
            "analyst",
            "decided_at",
        ]
    )
    for g in groups:
        dec = _latest_decision(g)
        asset = g.asset
        writer.writerow(
            [
                g.id,
                g.title,
                g.severity,
                g.tool,
                g.rule_id,
                asset.hostname if asset else "",
                asset.ip_address if asset else "",
                asset.port if asset else "",
                asset.protocol if asset else "",
                dec.decision.value if dec else g.queue_status.value,
                dec.reason if dec else "",
                dec.analyst if dec else "",
                dec.created_at.isoformat() if dec else "",
            ]
        )
    return buf.getvalue().encode("utf-8")


def export_tracker_xlsx(db: Session, engagement_id: int) -> bytes:
    groups = _confirmed_groups(db, engagement_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Confirmed Findings"
    headers = [
        "Finding Group ID",
        "Title",
        "Severity",
        "Tool",
        "Rule ID",
        "Hostname",
        "IP",
        "Port",
        "Protocol",
        "Decision",
        "Reason",
        "Analyst",
        "Decided At",
    ]
    ws.append(headers)
    for g in groups:
        dec = _latest_decision(g)
        asset = g.asset
        ws.append(
            [
                g.id,
                g.title,
                g.severity,
                g.tool,
                g.rule_id,
                asset.hostname if asset else None,
                asset.ip_address if asset else None,
                asset.port if asset else None,
                asset.protocol if asset else None,
                dec.decision.value if dec else g.queue_status.value,
                dec.reason if dec else None,
                dec.analyst if dec else None,
                dec.created_at.isoformat() if dec else None,
            ]
        )
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def export_report_docx(db: Session, engagement_id: int) -> bytes:
    engagement = db.get(Engagement, engagement_id)
    groups = _confirmed_groups(db, engagement_id)
    doc = Document()
    doc.add_heading("VAPT Findings Report", level=0)
    doc.add_paragraph(f"Engagement: {engagement.name if engagement else engagement_id}")
    if engagement and engagement.client:
        doc.add_paragraph(f"Client: {engagement.client}")
    doc.add_paragraph(
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    doc.add_paragraph(
        "This report includes only analyst-confirmed findings. "
        "Triage recommendations never auto-confirm or suppress findings."
    )
    doc.add_heading("Confirmed Findings", level=1)

    if not groups:
        doc.add_paragraph("No confirmed findings for this engagement.")
    else:
        for idx, g in enumerate(groups, start=1):
            dec = _latest_decision(g)
            asset = g.asset
            target = ""
            if asset:
                parts = [
                    p
                    for p in [
                        asset.hostname,
                        asset.ip_address,
                        f"{asset.port}/{asset.protocol}" if asset.port else None,
                    ]
                    if p
                ]
                target = " / ".join(str(p) for p in parts)
            doc.add_heading(f"{idx}. {g.title or g.rule_id}", level=2)
            doc.add_paragraph(f"Severity: {g.severity or 'n/a'}")
            doc.add_paragraph(f"Tool / Rule: {g.tool} / {g.rule_id}")
            doc.add_paragraph(f"Asset: {target or 'n/a'}")
            if dec:
                doc.add_paragraph(f"Analyst decision: {dec.decision.value}")
                if dec.reason:
                    doc.add_paragraph(f"Reason: {dec.reason}")
                if dec.analyst:
                    doc.add_paragraph(f"Analyst: {dec.analyst}")

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def write_export(path: Path, content: bytes) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    return path
